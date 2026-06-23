from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree

from polish_wlr_existing_docx import (
    A4_CONTENT_DXA,
    BLUE,
    DARK,
    MID_BLUE,
    MUTED,
    add_body,
    add_bullet,
    add_heading,
    configure_sections,
    configure_styles,
    create_table,
    format_table,
    move_elements_after,
    move_elements_before,
    rewrite_known_content,
    set_para_text,
    set_run_font,
    set_spacing,
)


ROOT = Path("/Users/yuxuan/Documents/轮足机器人")
SRC = ROOT / "outputs/docx_review/original_for_merge.docx"
OUT = ROOT / "outputs/docx_review/轮足巡检机器人技术方案与报价-整合修改意见版.docx"


def find_para(doc, exact):
    for p in doc.paragraphs:
        if p.text.strip() == exact:
            return p
    return None


def set_cell(cell, text):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_spacing(p, after=0, line=1.18)
    run = p.add_run(text)
    set_run_font(run, size=9.5, color=DARK)


def add_cover_stage_note(doc):
    subtitle = find_para(doc, "面向物流仓储 / 园区巡检场景")
    if subtitle is None:
        return
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(note, before=0, after=12, line=1)
    run = note.add_run("第一阶段样机研发版 · 已整合优化完善建议")
    set_run_font(run, size=11, color=MUTED, bold=True)
    subtitle._p.addnext(note._p)


def remove_existing_toc(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sdt"):
            continue
        xml_text = etree.tostring(child, encoding="unicode")
        if "Table of Contents" in xml_text or "TOC \\o" in xml_text or "目录" in xml_text:
            body.remove(child)


def clean_empty_toc_bookmark_paragraphs(doc):
    for p in list(doc.paragraphs):
        if p.text.strip():
            continue
        xml_text = etree.tostring(p._p, encoding="unicode")
        if "_Toc" not in xml_text:
            continue
        if "w:sectPr" in xml_text:
            p._p.getparent().remove(p._p)
            continue
        for child in list(p._p):
            if child.tag in {qn("w:bookmarkStart"), qn("w:bookmarkEnd")}:
                p._p.remove(child)
        p.style = "Normal"
        set_spacing(p, before=0, after=0, line=1)


def insert_manual_toc(doc):
    target = find_para(doc, "一、项目背景与目标")
    if target is None:
        return
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(title, before=0, after=8, line=1)
    run = title.add_run("目录")
    set_run_font(run, size=16, color=DARK, bold=True)
    toc = create_table(
        doc,
        ["章节", "本版整合重点"],
        [
            ["文档说明", "明确本版为第一阶段样机研发版，已整合优化完善建议。"],
            ["一、项目背景与目标", "新增研发阶段划分、本阶段边界和样机阶段不纳入范围。"],
            ["二、需求理解与技术路线选择", "保留原需求理解，强化样机研发的范围控制。"],
            ["三、总体系统架构", "保留端 - 边 - 云架构，作为后续系统扩展基础。"],
            ["四、硬件方案", "新增三套硬件采购规划、单套成本和三套合计预算。"],
            ["五、软件与算法二次开发方案", "新增四子系统解耦架构：巡检、MQTT 传输、流媒体、打点系统。"],
            ["六、样机研发计划与阶段交付物", "新增研发排期、硬件交付、软件交付和测试验收标准。"],
            ["七、商务报价", "将原空白报价表替换为样机阶段预算表，并保留后续量产参考价。"],
        ],
        [2600, 6426],
    )
    move_elements_before([title, toc], target)


def insert_stage_boundary(doc):
    target = find_para(doc, "二、需求理解与技术路线选择") or find_para(doc, "二、需求理解")
    if target is None:
        return
    elements = []
    elements.append(add_heading(doc, "1.3 研发阶段划分与本阶段边界", 2))
    elements.append(
        add_body(
            doc,
            "根据优化完善建议，本项目不再将量产落地、集群部署和大模型演进作为当前阶段直接交付目标，而采用“样机研发—验证机优化—量产机落地”的分层迭代路径。当前合同阶段聚焦第一阶段样机研发，核心任务是验证技术可行性、打通软硬件基础链路并输出可测试样机。",
        )
    )
    elements.append(
        create_table(
            doc,
            ["阶段", "定位", "目标边界", "核心交付"],
            [
                ["第一阶段", "样机研发", "完成单台机器人定点巡检、数据采集、视频回传、基础告警功能验证；不追求产品美观、全场景适配和规模化运维。", "主样机、测试样机、备件体系、V1.0 软件、样机测试报告。"],
                ["第二阶段", "验证机优化", "基于样机测试问题清单，优化外观、稳定性、操作流畅度、传感器精度和部署效率。", "验证机、优化版软件、现场试运行报告。"],
                ["第三阶段", "量产机落地", "开展多环境、长周期可靠性测试，固化硬件选型、软件版本和生产工艺。", "量产定型方案、生产测试规范、批量部署条件。"],
            ],
            [1450, 1700, 3926, 1950],
        )
    )
    elements.append(add_heading(doc, "1.4 样机阶段不纳入范围", 2))
    for item in [
        "不落地多机集群调度，仅预留后续调度系统接口。",
        "不落地 AI 缺陷识别、端到端 VLA 大模型训练和模型主控接管，仅预留数据采集与接口字段。",
        "不承诺全场景稳定运行和量产外观效果，样机阶段以功能链路打通、问题暴露和测试报告为验收重点。",
    ]:
        elements.append(add_bullet(doc, item))
    move_elements_before(elements, target)


def insert_hardware_budget(doc):
    target = find_para(doc, "五、软件与算法二次开发方案")
    if target is None:
        return
    elements = []
    elements.append(add_heading(doc, "4.4 三套硬件采购与研发预算", 2))
    elements.append(
        add_body(
            doc,
            "为避免样机研发过程中因单设备损坏、标定无对照样本或关键器件缺失导致研发停滞，硬件体系按三套配置规划：主样机、测试样机、备品备件。",
        )
    )
    elements.append(
        create_table(
            doc,
            ["硬件套件", "用途", "配置口径", "预算说明"],
            [
                ["第 1 套：主样机", "用于整机组装、基础巡检功能落地、客户演示和主流程联调。", "轮足本体、边缘计算、机械臂、传感器、相机、通信与结构件全套。", "按完整样机配置采购。"],
                ["第 2 套：测试样机", "用于并行功能测试、多传感器标定对比、算法效果验证和长时间压力测试。", "与主样机保持同等核心配置，便于数据对比和问题复现。", "按完整样机配置采购。"],
                ["第 3 套：备品备件", "覆盖机器狗本体、激光雷达、深度相机、机械臂、边缘计算模块等关键器件。", "以核心部件备份为主，允许部分外观件、非关键附件简配。", "按关键备件包采购。"],
            ],
            [1700, 3000, 2950, 1376],
        )
    )
    elements.append(add_heading(doc, "4.5 样机硬件成本暂估", 2))
    elements.append(
        create_table(
            doc,
            ["硬件项目", "配置说明", "不含税（万元）", "含税 13%（万元）"],
            [
                ["轮足机器人本体", "含基础传感器、电池、充电器、底盘 SDK。", "5.80", "6.55"],
                ["边缘计算模块", "Jetson Orin 系列、SSD、散热与电源转换。", "0.65", "0.73"],
                ["机械臂模组", "2~3 DoF 轻量机械臂、控制器、快换接口。", "0.80", "0.90"],
                ["气体与环境传感器", "CO / CO₂ / VOC / PM2.5 / 可燃气体等。", "0.60", "0.68"],
                ["视频与云台组件", "1080P 云台相机、编码与安装件。", "0.25", "0.28"],
                ["通信与结构辅件", "线束、支架、防护件、通信模块、调试辅件。", "0.35", "0.40"],
                ["手持终端与调试工具", "加固平板或 H5 终端、调试线材。", "0.30", "0.34"],
                ["单套小计", "完整样机硬件暂估。", "8.75", "9.89"],
            ],
            [2100, 3726, 1600, 1600],
        )
    )
    elements.append(
        create_table(
            doc,
            ["套件", "采购口径", "不含税（万元）", "含税 13%（万元）"],
            [
                ["主样机", "完整样机配置。", "8.75", "9.89"],
                ["测试样机", "完整样机配置。", "8.75", "9.89"],
                ["备品备件", "关键器件备件包，按完整配置约 80% 暂估。", "7.00", "7.91"],
                ["合计", "三套硬件研发采购预算。", "24.50", "27.69"],
            ],
            [1900, 3926, 1600, 1600],
        )
    )
    elements.append(add_body(doc, "注：以上预算为样机研发阶段暂估，最终金额应随底盘品牌、传感器型号、机械臂规格、是否采购备用电池及发票税率确认后更新。"))
    move_elements_before(elements, target)


def insert_software_decoupling(doc):
    target = find_para(doc, "5.2 端侧感知与决策算法")
    if target is None:
        return
    elements = []
    elements.append(add_heading(doc, "5.1.1 四子系统解耦架构（按修改意见新增）", 3))
    elements.append(
        add_body(
            doc,
            "为解决原方案中巡检业务、数据传输、视频流媒体和点位采集功能耦合的问题，样机阶段软件按“四个独立部署、接口协同”的方式重构。各子系统资源隔离、职责清晰，便于并行研发、故障定位和后续扩展。",
        )
    )
    elements.append(
        create_table(
            doc,
            ["子系统", "核心职责", "样机阶段实现范围", "关键接口"],
            [
                ["独立巡检系统", "承载核心业务流程：点位、任务、路径、告警、记录、导出。", "点位管理、自动巡检、基础路径规划、异常告警、历史查询。", "REST / WebSocket / 数据库接口。"],
                ["传感器采集与传输系统", "采集气体、视觉、设备姿态等数据，完成边缘预处理与云端传输。", "阿里云 IoT 平台 MQTT 双向通信、离线缓存、断点续传。", "MQTT Topic / 标准数据 Schema。"],
                ["独立流媒体服务器", "承载视频实时回传、录制、回放和远程查看。", "WebRTC 低延迟方案，支持 2~4 路 1080P 视频。", "WebRTC / RTSP / 录像索引。"],
                ["独立打点系统", "作为部署研发工具，采集世界坐标、机械臂姿态、相机参数并生成配置。", "人工辅助精准打点、配置文件导出导入、批量点位管理。", "视频画面、MQTT 数据、点位配置文件。"],
            ],
            [1850, 2776, 2900, 1500],
        )
    )
    move_elements_before(elements, target)


def insert_detailed_subsystems(doc):
    target = find_para(doc, "七、商务报价") or find_para(doc, "六、商务报价")
    if target is None:
        return
    elements = []
    elements.append(add_heading(doc, "5.5 独立软件子系统建设内容", 2))
    elements.append(add_heading(doc, "5.5.1 独立巡检系统", 3))
    for item in [
        "输出需求调研文档、业务逻辑文档，明确园区巡检点位、检查内容、告警规则和任务流程。",
        "完成数据库设计，分层存储巡检任务、点位信息、检测数据、告警日志和设备状态。",
        "迭代前端原型与正式前后端页面，实现点位管理、路径规划、任务调度、实时监控、历史回放和数据导出。",
        "样机阶段仅实现基础业务闭环，高阶排班、跨园区运营和复杂调度延后至后续阶段。",
    ]:
        elements.append(add_bullet(doc, item))
    elements.append(add_heading(doc, "5.5.2 独立传感器信息采集与传输系统", 3))
    elements.append(
        add_body(
            doc,
            "传感器采集与传输系统基于阿里云 IoT 平台 MQTT 消息服务器实现设备与云端双向通信。所有传感器数据经边缘计算预处理后，通过标准 Topic 上传，独立于巡检业务逻辑，仅通过标准接口与巡检系统协同。",
        )
    )
    for item in [
        "采集范围包括气体浓度、环境数据、图像元数据、设备姿态、运行状态和告警事件。",
        "弱网策略包括本地缓存、断点续传、失败重试和上传状态追踪，避免园区弱网环境下数据丢失。",
        "数据规范统一时间戳、设备编号、点位编号、任务编号和数据类型，便于后续查询与训练数据积累。",
    ]:
        elements.append(add_bullet(doc, item))
    elements.append(add_heading(doc, "5.5.3 独立流媒体服务器", 3))
    for item in [
        "技术选型采用 WebRTC 低延迟方案，支持 2~4 路 1080P 视频并发回传、录制和回放。",
        "部署带 GPU 转码能力的专属节点，不与数据库、MQTT 服务共用关键资源，避免视频高占用干扰巡检业务。",
        "建议园区上行带宽 ≥ 20 Mbps，保障远程遥控、点位打点和实时监控流畅性。",
    ]:
        elements.append(add_bullet(doc, item))
    elements.append(add_heading(doc, "5.5.4 独立打点系统（专属研发工具）", 3))
    for item in [
        "采集机器人世界坐标、机械臂姿态、相机参数、巡检点位名称和点位检查内容。",
        "对接流媒体视频画面与 MQTT 传感器数据，辅助人工精准打点。",
        "生成标准化点位配置文件并批量导入巡检系统，避免运行时巡检系统承担高并发点位采集压力。",
    ]:
        elements.append(add_bullet(doc, item))
    move_elements_before(elements, target)


def insert_plan_and_acceptance(doc):
    target = find_para(doc, "七、商务报价") or find_para(doc, "六、商务报价")
    if target is None:
        return
    elements = []
    elements.append(add_heading(doc, "六、样机研发计划与阶段交付物", 1))
    elements.append(
        create_table(
            doc,
            ["阶段", "周期", "研发任务", "交付物"],
            [
                ["需求与架构冻结", "第 1-2 周", "现场需求梳理、点位样例确认、四子系统接口定义、硬件采购清单冻结。", "需求调研文档、业务逻辑文档、数据库草案、接口清单。"],
                ["硬件组装与底层联调", "第 3-5 周", "三套硬件到货、主样机装配、传感器接入、机械臂动作模板、网络链路测试。", "主样机、测试样机、备件清单、硬件调试记录。"],
                ["软件 V1.0 开发", "第 6-9 周", "巡检系统、MQTT 传输系统、流媒体服务、打点工具并行开发。", "四个独立子系统 V1.0、数据库结构、部署脚本。"],
                ["集成测试与问题闭环", "第 10-12 周", "整机巡检流程联调、弱网测试、视频回传测试、点位导入导出测试。", "样机测试报告、问题清单、整改建议、验证机阶段规划。"],
            ],
            [1600, 1400, 3726, 2300],
        )
    )
    elements.append(add_heading(doc, "6.1 硬件交付", 2))
    elements.append(add_body(doc, "完成三套硬件装配调试，形成主样机、测试样机和备品备件全套硬件体系。交付时应确认整机功能正常、传感器标定完成、通信链路通畅。"))
    elements.append(add_heading(doc, "6.2 软件交付（V1.0 样机版本）", 2))
    elements.append(
        create_table(
            doc,
            ["交付项", "样机阶段功能"],
            [
                ["独立巡检系统", "点位管理、自动巡检、路径规划、异常告警、数据查询导出基础功能闭环。"],
                ["传感器传输系统", "MQTT 双向通信、数据预处理、离线缓存、断点续传稳定运行。"],
                ["流媒体服务器", "支持多路 1080P 视频实时回传、录制、低延迟远程查看。"],
                ["独立打点工具", "支持园区点位精准采集、配置文件导出导入，满足部署需求。"],
            ],
            [2600, 6426],
        )
    )
    elements.append(add_heading(doc, "6.3 测试验收标准", 2))
    elements.append(
        create_table(
            doc,
            ["验收类别", "样机阶段验收标准", "验收方式"],
            [
                ["基础运动", "机器人可完成前进、后退、转向、停止、遥控接管和急停。", "现场控制测试与异常停机测试。"],
                ["定点巡检", "可按导入点位执行巡检任务，到点后触发拍照、气体采集和告警判断。", "不少于 10 个点位连续执行 3 轮。"],
                ["传感器传输", "气体、姿态、设备状态等数据可通过 MQTT 上传，弱网后可补传。", "断网 / 恢复网络模拟测试。"],
                ["视频系统", "支持 2~4 路 1080P 视频回传、录制和回放，远程查看延迟满足样机验证。", "局域网与园区网络环境测试。"],
                ["打点系统", "可采集坐标、机械臂姿态和相机参数，并生成可导入巡检系统的配置文件。", "人工打点、导出、导入、执行闭环测试。"],
                ["文档交付", "需求、数据库、接口、部署、测试报告和问题清单完整。", "文档审查与交付清单核对。"],
            ],
            [1850, 4576, 2600],
        )
    )
    move_elements_before(elements, target)


def update_budget_tables(doc):
    budget_rows = [
        ["三套硬件采购", "主样机、测试样机、备品备件，含关键传感器、机械臂、边缘计算与结构辅件。", "约 27.69"],
        ["需求与架构设计", "需求调研、业务逻辑、数据库、接口、部署架构和技术文档。", "8 ~ 12"],
        ["机器人端集成", "底盘接口、传感器接入、机械臂动作模板、日志和状态采集。", "15 ~ 25"],
        ["独立巡检系统", "点位、任务、路径、告警、查询导出和基础前后端页面。", "20 ~ 30"],
        ["MQTT 传输系统", "阿里云 IoT 接入、数据预处理、离线缓存、断点续传。", "12 ~ 18"],
        ["独立流媒体服务器", "WebRTC 视频回传、录制、回放和专属节点部署。", "15 ~ 22"],
        ["独立打点系统", "坐标、姿态、相机参数采集，点位配置导入导出。", "10 ~ 15"],
        ["联调测试与交付", "现场联调、测试报告、问题清单、培训与验收支持。", "10 ~ 15"],
        ["样机阶段合计", "第一阶段样机研发总体预算。", "约 117.69 ~ 164.69"],
    ]
    for table in doc.tables:
        first = table.cell(0, 0).text.strip()
        if first == "研发模块" and len(table.rows) >= 10 and len(table.columns) >= 3:
            set_cell(table.cell(0, 0), "预算模块")
            set_cell(table.cell(0, 1), "主要工作量")
            set_cell(table.cell(0, 2), "费用（万元，含税暂估）")
            for idx, row_data in enumerate(budget_rows, start=1):
                for c_idx, value in enumerate(row_data):
                    set_cell(table.cell(idx, c_idx), value)
        if first == "适用阶段" and len(table.rows) >= 5:
            set_cell(table.cell(0, 0), "后续适用阶段")
            set_cell(table.cell(0, 1), "单台出厂价（参考）")
            rows = [
                ["验证机 / 小批量", "9.8 万元"],
                ["初次规模采购", "9.2 万元"],
                ["市场铺开", "8.6 万元"],
                ["批量量产", "≤ 7.8 万元"],
            ]
            for idx, row_data in enumerate(rows, start=1):
                for c_idx, value in enumerate(row_data):
                    set_cell(table.cell(idx, c_idx), value)

    intro = find_para(doc, "本项目报价采用“一次性研发费（NRE）+ 量产整机单价 + 软件年度服务费”的三段式结构。该结构能保证整机出厂价控制在 10 万元以内，同时为研发投入回收与长期服务收益建立合理路径。")
    if intro is not None:
        set_para_text(
            intro,
            "本项目报价按“样机研发服务费 + 三套硬件采购费 + 后续阶段参考费用”拆分。当前阶段以样机研发预算为主，验证机、量产机和远期平台能力不提前混入样机阶段交付范围。",
        )

    note = find_para(doc, "说明：以上均为含税出厂价（增值税专票 13%）；累计订购量可按 12 个月滚动计算，特殊定制（防爆、低温、特定气体探头等）单独议价。")
    if note is not None:
        set_para_text(note, "说明：后续单台出厂价仅作为验证机 / 量产机阶段参考，最终报价需依据样机测试结果、量产 BOM、生产工艺和采购规模重新核定。")


def add_post_budget_note(doc):
    target = find_para(doc, "7.3 软件年度服务费")
    if target is None:
        return
    elements = [
        add_body(doc, "样机阶段以交付 V1.0 可测试系统为主，不强制绑定年度运维服务费。若后续进入验证机或量产机阶段，可再按设备数量、私有化部署范围和服务等级确认年度服务费。")
    ]
    move_elements_after(elements, target)


def main():
    doc = Document(SRC)
    configure_styles(doc)
    configure_sections(doc)
    rewrite_known_content(doc)
    remove_existing_toc(doc)
    clean_empty_toc_bookmark_paragraphs(doc)
    add_cover_stage_note(doc)
    insert_stage_boundary(doc)
    insert_hardware_budget(doc)
    insert_software_decoupling(doc)
    insert_detailed_subsystems(doc)
    insert_plan_and_acceptance(doc)
    update_budget_tables(doc)
    add_post_budget_note(doc)
    insert_manual_toc(doc)
    for table in doc.tables:
        format_table(table)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"OK: {OUT}")


if __name__ == "__main__":
    main()
