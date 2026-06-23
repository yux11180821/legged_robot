from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path("/Users/yuxuan/Documents/轮足机器人")
SRC = ROOT / "outputs/docx_review/source.docx"
OUT = ROOT / "outputs/docx_review/轮足巡检机器人技术方案与报价-精修版.docx"

FONT = "Microsoft YaHei"
BLUE = "1F4E78"
MID_BLUE = "2E74B5"
DARK = "1F2933"
MUTED = "667085"
LIGHT_BLUE = "EAF3F8"
LIGHT_GRAY = "F3F5F7"
GRID = "C9D1D9"
A4_CONTENT_DXA = 9026
TABLE_INDENT_DXA = 0


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_paragraph(p):
    for child in list(p._p):
        if child.tag == qn("w:r"):
            p._p.remove(child)


def set_para_text(p, text, size=10.5, color=DARK, bold=False, italic=False):
    clear_paragraph(p)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return run


def remove_paragraph(p):
    parent = p._element.getparent()
    parent.remove(p._element)


def paragraph_border_bottom(p, color=BLUE, size=12, space=8):
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)


def set_spacing(p, before=0, after=6, line=1.25):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def apply_bullet(p, num_id="1", level="0"):
    try:
        p.style = "List Paragraph"
    except KeyError:
        p.style = "Normal"
    p_pr = p._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), level)
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), num_id)
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, MID_BLUE, 12, 7),
        ("Heading 3", 11.5, BLUE, 8, 5),
    ]:
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.18

    for style_name in ["List Bullet", "List Number", "List Paragraph"]:
        if style_name in styles:
            style = styles[style_name]
            style.font.name = FONT
            style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            style.font.size = Pt(10.5)
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.line_spacing = 1.22


def configure_sections(doc):
    for section in doc.sections:
        section.page_width = Twips(11906)
        section.page_height = Twips(16838)
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.45)
        section.footer_distance = Inches(0.4)
        clear_part(section.header)
        clear_part(section.footer)

        hp = section.header.add_paragraph()
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_spacing(hp, after=2, line=1)
        hr = hp.add_run("轮足巡检机器人项目 · 技术方案与商务报价")
        set_run_font(hr, size=8.5, color=MUTED)
        paragraph_border_bottom(hp, color="D0D5DD", size=4, space=4)

        fp = section.footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(fp, after=0, line=1)
        r = fp.add_run("内部讨论稿 · 以正式合同与技术协议为准")
        set_run_font(r, size=8.5, color=MUTED)


def clear_part(part):
    for paragraph in list(part.paragraphs):
        remove_paragraph(paragraph)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")


def set_cell_margins(cell, top=100, bottom=100, start=140, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_tc_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Twips(width)


def set_table_geometry(table, widths, indent=TABLE_INDENT_DXA):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl_pr.addnext(grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_tc_width(cell, widths[min(idx, len(widths) - 1)])


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), GRID)


def format_table(table):
    widths_by_cols = {
        1: [A4_CONTENT_DXA],
        2: [2350, A4_CONTENT_DXA - 2350],
        3: [2050, 4900, A4_CONTENT_DXA - 6950],
        4: [1700, 2750, 2600, A4_CONTENT_DXA - 7050],
    }
    widths = widths_by_cols.get(len(table.columns), [A4_CONTENT_DXA // len(table.columns)] * len(table.columns))
    set_table_geometry(table, widths)
    set_table_borders(table)
    for r_idx, row in enumerate(table.rows):
        if r_idx == 0 and len(table.rows) > 1:
            tr_pr = row._tr.get_or_add_trPr()
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0 and len(table.rows) > 1:
                shade_cell(cell, BLUE)
            elif len(table.rows) == 1:
                shade_cell(cell, LIGHT_BLUE)
            elif r_idx % 2 == 0:
                shade_cell(cell, LIGHT_GRAY)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (r_idx == 0 or is_short_column(table, c_idx)) else WD_ALIGN_PARAGRAPH.LEFT
                set_spacing(p, after=0, line=1.18)
                for run in p.runs:
                    set_run_font(
                        run,
                        size=9.5 if len(table.columns) >= 3 else 10,
                        color="FFFFFF" if r_idx == 0 and len(table.rows) > 1 else DARK,
                        bold=(r_idx == 0 or c_idx == 0),
                    )


def is_short_column(table, col_idx):
    header = table.cell(0, col_idx).text.strip()
    short_tokens = ["数量", "费用", "单价", "时间", "阶段", "版本", "出厂价"]
    return any(token in header for token in short_tokens) or col_idx == len(table.columns) - 1 and len(table.columns) >= 3


def create_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    if widths:
        set_table_geometry(table, widths)
    format_table(table)
    return table


def append_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 13, 3: 11.5}[level], color=BLUE if level != 2 else MID_BLUE, bold=True)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_spacing(p, after=6, line=1.28)
    for run in p.runs:
        set_run_font(run, size=10.5, color=DARK)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(text)
    apply_bullet(p)
    set_spacing(p, after=4, line=1.22)
    for run in p.runs:
        set_run_font(run, size=10.5, color=DARK)
    return p


def move_elements_before(elements, target_p):
    for element in elements:
        target_p._p.addprevious(element._element if hasattr(element, "_element") else element._tbl)


def move_elements_after(elements, target_p):
    anchor = target_p._p
    for element in elements:
        xml = element._element if hasattr(element, "_element") else element._tbl
        anchor.addnext(xml)
        anchor = xml


def rewrite_known_content(doc):
    replacements = {
        "二、需求理解": "二、需求理解与技术路线选择",
        "四、硬件方案": "四、硬件方案",
        "基于Jetson Orin系列嵌入式GPU平台，部署我方感知与决策算法。": "基于 Jetson Orin 系列嵌入式 GPU 平台，部署我方感知与决策算法。",
        "5.2.2 巡检点识别与辅助定位": "5.2.2 巡检点识别与精确对位",
        "本项目首期只做摄像拍照，不做图像识别和缺陷识别。系统支持人工拍照、到点自动拍照和任务触发拍照。": "本项目首期聚焦图像采集、留存与质量控制，暂不纳入复杂缺陷识别。系统支持人工拍照、到点自动拍照和任务触发拍照。",
        "机械臂机械臂不做连续轨迹规划和高精度力控，确保系统简单可靠，控制功能包括：": "机械臂首期不做连续轨迹规划和高精度力控，优先保证系统简单可靠。控制功能包括：",
        "5.3.2 本地缓存与断点续传": "5.3.3 本地缓存与断点续传",
        "5.3.3 上位机功能": "5.3.4 上位机功能",
        "5.3.4 手持终端功能": "5.3.5 手持终端功能",
        "5.4 软件交付内容清单": "5.4 软件交付内容清单",
        "六、商务报价": "七、商务报价",
        "6.1 一次性研发费": "7.1 一次性研发费",
        "6.2 量产整机单价": "7.2 量产整机单价",
        "6.3 运维服务费": "7.3 软件年度服务费",
        "说明：以上均为含税出厂价（增值税专票___%）。": "说明：以上均为含税出厂价（增值税专票 13%）；累计订购量可按 12 个月滚动计算，特殊定制（防爆、低温、特定气体探头等）单独议价。",
    }
    delete_exact = {"巡检点识别与精确对位"}

    h1_texts = {"文档说明", "一、项目背景与目标", "二、需求理解与技术路线选择", "三、总体系统架构", "四、硬件方案", "五、软件与算法二次开发方案", "七、商务报价"}
    h2_pattern = re.compile(r"^\d+\.\d+ ")
    h3_pattern = re.compile(r"^\d+\.\d+\.\d+ ")

    list_like = {
        "读取机器人当前位置、速度、运行状态；",
        "接入 2D 地图或导航点位信息；",
        "若接口支持，采集 3D 点云或深度感知数据；",
        "记录机器人在巡检过程中的位姿轨迹；",
        "将图片、气体数据、位姿和任务状态按统一时间戳保存。",
        "多传感器时间同步；",
        "2D 栅格地图与巡检点位绑定；",
        "3D 点云回放；",
        "障碍物语义标注；",
        "用于 VLA 训练的数据对齐与导出。",
        "手持终端一键拍照；",
        "巡检点到达后自动拍照；",
        "机械臂动作到位后拍照；",
        "图片本地保存；",
        "图片自动上传后台；",
        "图片按任务、点位、时间归档；",
        "网络中断时本地缓存，恢复后自动补传。",
        "原厂遥控器控制；",
        "Web 上位机控制；",
        "手持终端 H5 / App 控制；",
        "远程网络控制；",
        "急停控制。",
        "巡检图片；",
        "气体检测数据；",
        "巡检任务记录；",
        "告警记录；",
        "机器人状态日志；",
        "操作日志。",
        "查看机器人连接状态；",
        "查看电量和网络状态；",
        "手动遥控机器人；",
        "一键拍照；",
        "查看实时气体数据；",
        "查看告警提示；",
        "控制机械臂简单动作；",
        "启动 / 暂停巡检任务；",
        "查看最近巡检记录。",
        "异常时自动停止；",
        "任务结束后自动收回。",
        "巡检任务自动调用预设动作；",
        "动作执行完成后触发拍照或气体检测；",
    }

    for p in list(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        if text in delete_exact:
            remove_paragraph(p)
            continue
        new_text = replacements.get(text, text)

        if new_text != text:
            set_para_text(p, new_text)
            text = new_text

        if text in {"轮足巡检机器人项目", "技 术 方 案 与 商 务 报 价 书", "面向物流仓储 / 园区巡检场景"}:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_spacing(p, before=0, after=8, line=1.0)
            sizes = {
                "轮足巡检机器人项目": 26,
                "技 术 方 案 与 商 务 报 价 书": 20,
                "面向物流仓储 / 园区巡检场景": 12,
            }
            colors = {
                "轮足巡检机器人项目": BLUE,
                "技 术 方 案 与 商 务 报 价 书": BLUE,
                "面向物流仓储 / 园区巡检场景": MUTED,
            }
            for r in p.runs:
                set_run_font(r, size=sizes[text], color=colors[text], bold=text != "面向物流仓储 / 园区巡检场景")
        elif text in h1_texts:
            p.style = "Heading 1"
        elif text in {"4.1 轮足运动平台", "4.2 感知与作业扩展", "4.3 硬件交付内容清单", "5.4 软件交付内容清单", "7.1 一次性研发费", "7.2 量产整机单价", "7.3 软件年度服务费"} or h2_pattern.match(text):
            p.style = "Heading 2"
        elif h3_pattern.match(text) or text in {"端：机器人本体层", "边：园区 / 仓库边缘服务器层", "云：跨园区 / 跨客户数据中台层"}:
            p.style = "Heading 3"
        elif text in list_like or re.match(r"^（\\d+）", text):
            if re.match(r"^（\\d+）", text):
                text = re.sub(r"^（\\d+）\\s*", "", text)
                set_para_text(p, text)
            apply_bullet(p)
        elif p.style and p.style.name == "Normal (Web)":
            p.style = "Normal"

        if p.style and p.style.name.startswith("Heading"):
            for r in p.runs:
                level = int(p.style.name.split()[-1])
                set_run_font(r, size={1: 16, 2: 13, 3: 11.5}.get(level, 10.5), color=BLUE if level != 2 else MID_BLUE, bold=True)
        elif p.style and p.style.name in {"List Bullet", "List Paragraph"}:
            set_spacing(p, after=4, line=1.22)
            for r in p.runs:
                set_run_font(r, size=10.5, color=DARK)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_spacing(p, after=6, line=1.28)
            for r in p.runs:
                set_run_font(r, size=10.5, color=DARK)


def polish_cover(doc):
    subtitle = next((p for p in doc.paragraphs if p.text.strip() == "面向物流仓储 / 园区巡检场景"), None)
    if subtitle is None:
        return
    meta = doc.add_table(rows=3, cols=2)
    rows = [
        ("文档类型", "技术方案与商务报价书"),
        ("项目场景", "物流仓储 / 园区巡检"),
        ("版本日期", "精修版 · 2026年5月"),
    ]
    for i, row in enumerate(rows):
        meta.rows[i].cells[0].text = row[0]
        meta.rows[i].cells[1].text = row[1]
    format_table(meta)
    subtitle._p.addnext(meta._tbl)

    rule = doc.add_paragraph()
    paragraph_border_bottom(rule, color=BLUE, size=10, space=8)
    meta._tbl.addnext(rule._p)


def fill_quote_tables(doc):
    nre_rows = [
        ("硬件集成与结构设计", "25"),
        ("端侧感知与决策算法", "35"),
        ("端侧应用任务层", "20"),
        ("上位机与手持终端", "25"),
        ("云端数据中台", "35"),
        ("集群调度引擎", "20"),
        ("试点联调与产品化", "15"),
        ("项目管理与质量", "10"),
        ("NRE 合计", "185"),
    ]
    for table in doc.tables:
        if table.cell(0, 0).text.strip() == "研发模块":
            for row in table.rows[1:]:
                module = row.cells[0].text.strip()
                for key, value in nre_rows:
                    if module == key:
                        row.cells[2].text = value
                        break
        if table.cell(0, 0).text.strip() == "适用阶段":
            values = [
                ("试点 / 小批量（1~9 台）", "9.8 万元"),
                ("初次规模采购（10~49 台）", "9.2 万元"),
                ("市场铺开（50~199 台）", "8.6 万元"),
                ("批量量产（≥200 台）", "≤ 7.8 万元"),
            ]
            for row, (stage, price) in zip(table.rows[1:], values):
                row.cells[0].text = stage
                row.cells[1].text = price


def add_project_plan(doc):
    target = next((p for p in doc.paragraphs if p.text.strip() == "七、商务报价"), None)
    if target is None:
        return
    elements = []
    elements.append(add_heading(doc, "六、项目实施计划与交付物", 1))
    elements.append(add_body(doc, "项目建议按“原型验证—试制联调—现场试点—批量交付”四阶段推进，确保技术风险在小批量前充分暴露并闭环。"))
    elements.append(
        create_table(
            doc,
            ["阶段", "周期", "主要工作", "阶段交付物"],
            [
                ["原型验证", "第 1-2 月", "需求冻结、底盘与传感器选型、原型机改装、端侧软件框架搭建。", "原型机 1 台、详细设计文档、接口清单。"],
                ["试制联调", "第 3-4 月", "端侧任务执行、上位机、手持终端、数据入库与基础告警闭环。", "试制机 3~5 台、上位机 V1.0、终端应用 V1.0。"],
                ["现场试点", "第 5 月", "甲方指定仓库 / 园区现场部署，连续运行测试与问题闭环。", "试点验收报告、问题闭环清单、用户手册。"],
                ["批量交付", "第 6 月起", "量产备料、出厂检测、培训交付、运维 SOP 落地。", "批量整机、培训记录、运维手册。"],
            ],
            [1600, 1500, 3350, 2576],
        )
    )
    elements.append(add_heading(doc, "6.1 交付物范围", 2))
    for text in [
        "整机硬件：轮足机器人本体、传感器扩展、机械臂模组、充电与换电配件、手持终端。",
        "机器人端软件：底盘接口封装、巡检任务执行、拍照 / 气体 / 机械臂控制、日志记录与断点续传。",
        "管理端软件：Web 上位机、任务配置、设备状态、告警中心、巡检记录与数据导出。",
        "文档与培训：部署手册、用户手册、运维 SOP、接口文档和现场培训材料。",
    ]:
        elements.append(add_bullet(doc, text))
    move_elements_before(elements, target)


def add_service_fee_table(doc):
    target = next((p for p in doc.paragraphs if p.text.strip() == "7.3 软件年度服务费"), None)
    if target is None:
        return
    elements = [
        add_body(doc, "从交付第二年起，建议按机器人台数收取年度软件服务费，覆盖中台服务、远程运维、版本升级与故障响应。"),
        create_table(
            doc,
            ["服务等级", "单台年费", "服务范围"],
            [
                ["基础版", "8,000 元 / 年", "SaaS 中台、OTA、5×8 远程支持、年度 1 次升级。"],
                ["标准版", "15,000 元 / 年", "基础版 + 7×24 远程支持、季度升级、24 小时备件响应。"],
                ["企业版", "28,000 元 / 年", "标准版 + 私有化部署、按需现场支持、客户定制需求池。"],
            ],
            [2200, 2200, A4_CONTENT_DXA - 4400],
        ),
    ]
    move_elements_after(elements, target)


def add_tail_sections(doc):
    append_page_break(doc)
    add_heading(doc, "八、验收标准与测试方案", 1)
    add_body(doc, "建议采用“出厂验收 + 现场试点验收 + 批量抽检验收”的三级验收机制，避免停留在演示效果，确保产品具备真实交付能力。")
    create_table(
        doc,
        ["验收项", "建议标准", "测试方法"],
        [
            ["移动能力", "巡检速度 0.5~1.0 m/s；可稳定通过小坡道、托盘边缘等典型障碍。", "标准路线往返测试，统计通行成功率与异常停机次数。"],
            ["续航与换电", "单次续航 ≥ 2 h；人工换电 ≤ 1 min；低电量回桩策略正常触发。", "连续巡航、换电计时与低电量模拟。"],
            ["自主巡检", "按预设路线完成到点、拍照、气体检测、数据上传；漏检率 ≤ 2%。", "不少于 20 个点位连续执行 5 轮。"],
            ["安全避障", "动态人员靠近时减速 / 停障；急停链路触发时间 ≤ 300 ms。", "人员横穿、静态障碍、软硬件急停联动测试。"],
            ["数据闭环", "巡检图片、气体数据、任务日志与告警记录可检索、导出、追溯。", "抽样核对字段完整性、时间戳与点位绑定关系。"],
        ],
        [1850, 4076, 3100],
    )

    add_heading(doc, "九、关键风险与应对", 1)
    create_table(
        doc,
        ["风险类别", "风险描述", "应对策略"],
        [
            ["硬件平台供应", "底盘或关键传感器存在涨价、交期或断供风险。", "运动接口适配层与关键件双供应商预案并行设计。"],
            ["现场环境复杂", "仓库布局、光照、信号覆盖、地面障碍差异较大。", "试点期驻场调参，沉淀标准部署 SOP 与场地准入清单。"],
            ["软件范围膨胀", "中台、算法、终端需求不断扩展，影响交付周期。", "按首期必需 / 二期增强拆分范围，双周评审需求变更。"],
            ["人员安全", "人车混行环境下存在碰撞与误操作风险。", "三重急停、速度限幅、虚拟围栏、物理防撞与责任险组合。"],
            ["数据安全", "巡检图片与园区数据涉及客户敏感信息。", "支持私有化部署、权限隔离、日志审计与数据脱敏策略。"],
        ],
        [1800, 3300, 3926],
    )

    add_heading(doc, "十、后续合作与演进方向", 1)
    for text in [
        "基于真实巡检数据沉淀 VLA 训练数据集，逐步从规则式巡检演进到端到端任务执行。",
        "扩展低温、防爆、户外、消防侦察等特种型号，形成多场景轮足巡检产品矩阵。",
        "与甲方现有 WMS、园区安防、EHS 平台对接，形成统一告警与工单闭环。",
    ]:
        add_bullet(doc, text)


def main():
    doc = Document(SRC)
    configure_styles(doc)
    configure_sections(doc)
    rewrite_known_content(doc)
    polish_cover(doc)
    fill_quote_tables(doc)
    add_project_plan(doc)
    add_service_fee_table(doc)
    add_tail_sections(doc)
    for table in doc.tables:
        format_table(table)
    doc.save(OUT)
    print(f"OK: {OUT}")


if __name__ == "__main__":
    main()
